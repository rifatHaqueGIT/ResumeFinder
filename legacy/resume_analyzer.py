r"""
Resume Analyzer — reads each resume found by resume_finder.py, extracts content,
performs expert hiring-manager analysis, and creates a formatted Excel spreadsheet.

Produces:
  - Resume name & location
  - Tagged role
  - Company mentions
  - Hiring manager score (0-100)
  - Role-specific keyword checklist with FOUND / MISSING highlighting
"""

import csv
import os
import re
import sys
from pathlib import Path
from collections import defaultdict

# Force UTF-8 on Windows
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from PyPDF2 import PdfReader
    HAS_PYPDF2 = True
except ImportError:
    HAS_PYPDF2 = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ═══════════════════════════════════════════════════════════════════════════════
# ROLE DEFINITIONS — keywords an expert hiring manager looks for per role
# ═══════════════════════════════════════════════════════════════════════════════

ROLE_KEYWORDS = {
    "Software Engineer": {
        "Technical Skills": [
            "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust",
            "sql", "html", "css", "react", "angular", "vue", "node.js", "spring",
            "django", "flask", "fastapi", ".net", "rest api", "graphql",
        ],
        "DevOps & Tools": [
            "git", "docker", "kubernetes", "ci/cd", "jenkins", "github actions",
            "aws", "azure", "gcp", "terraform", "linux", "bash",
        ],
        "Core Competencies": [
            "data structures", "algorithms", "system design", "microservices",
            "object-oriented", "design patterns", "testing", "unit test",
            "agile", "scrum", "code review", "debugging",
        ],
        "Soft Skills & Impact": [
            "leadership", "teamwork", "collaboration", "communication",
            "problem solving", "mentoring", "cross-functional",
        ],
        "Experience Signals": [
            "internship", "co-op", "full-time", "contract",
            "project", "developed", "implemented", "built", "designed",
            "optimized", "reduced", "increased", "improved", "deployed",
        ],
    },
    "Software Developer": {  # alias, same keywords
        "Technical Skills": [
            "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust",
            "sql", "html", "css", "react", "angular", "vue", "node.js", "spring",
            "django", "flask", "fastapi", ".net", "rest api", "graphql",
        ],
        "DevOps & Tools": [
            "git", "docker", "kubernetes", "ci/cd", "jenkins", "github actions",
            "aws", "azure", "gcp", "terraform", "linux", "bash",
        ],
        "Core Competencies": [
            "data structures", "algorithms", "system design", "microservices",
            "object-oriented", "design patterns", "testing", "unit test",
            "agile", "scrum", "code review", "debugging",
        ],
        "Soft Skills & Impact": [
            "leadership", "teamwork", "collaboration", "communication",
            "problem solving", "mentoring", "cross-functional",
        ],
        "Experience Signals": [
            "internship", "co-op", "full-time", "contract",
            "project", "developed", "implemented", "built", "designed",
            "optimized", "reduced", "increased", "improved", "deployed",
        ],
    },
    "Web Developer": {
        "Frontend": [
            "html", "css", "javascript", "typescript", "react", "angular", "vue",
            "next.js", "sass", "tailwind", "bootstrap", "responsive design",
            "webpack", "vite", "figma",
        ],
        "Backend": [
            "node.js", "express", "django", "flask", "php", "ruby on rails",
            "rest api", "graphql", "database", "sql", "mongodb", "postgresql",
        ],
        "DevOps & Tools": [
            "git", "docker", "aws", "vercel", "netlify", "ci/cd",
            "testing", "jest", "cypress", "seo", "accessibility",
        ],
        "Soft Skills & Impact": [
            "collaboration", "communication", "agile", "scrum",
            "project", "portfolio", "deployed", "built", "designed",
        ],
    },
    "Data Scientist": {
        "Core ML/AI": [
            "machine learning", "deep learning", "neural network", "nlp",
            "computer vision", "reinforcement learning", "generative ai",
            "tensorflow", "pytorch", "keras", "scikit-learn", "xgboost",
        ],
        "Data & Analytics": [
            "python", "r", "sql", "pandas", "numpy", "matplotlib", "seaborn",
            "jupyter", "tableau", "power bi", "spark", "hadoop", "etl",
            "data pipeline", "feature engineering", "a/b testing",
        ],
        "Statistics & Math": [
            "statistics", "probability", "regression", "classification",
            "clustering", "dimensionality reduction", "hypothesis testing",
            "bayesian", "optimization",
        ],
        "Soft Skills & Impact": [
            "communication", "presentation", "stakeholder",
            "research", "published", "accuracy", "improved", "reduced",
        ],
    },
    "SDET": {
        "Testing Core": [
            "test automation", "selenium", "cypress", "playwright", "appium",
            "junit", "testng", "pytest", "jest", "mocha",
            "api testing", "postman", "rest assured",
            "performance testing", "load testing", "jmeter",
        ],
        "Development": [
            "python", "java", "javascript", "typescript", "c#",
            "git", "docker", "ci/cd", "jenkins", "github actions",
        ],
        "QA Practices": [
            "test plan", "test cases", "bug tracking", "jira",
            "regression", "integration testing", "unit testing",
            "tdd", "bdd", "cucumber", "quality assurance",
        ],
        "Soft Skills": [
            "agile", "scrum", "collaboration", "communication",
            "problem solving", "attention to detail",
        ],
    },
    "Product Manager": {
        "Core PM": [
            "product strategy", "roadmap", "user stories", "requirements",
            "stakeholder", "prioritization", "market research",
            "competitive analysis", "go-to-market", "product lifecycle",
        ],
        "Tools & Methods": [
            "jira", "confluence", "figma", "analytics", "a/b testing",
            "sql", "data-driven", "kpi", "okr", "agile", "scrum",
        ],
        "Soft Skills": [
            "leadership", "communication", "cross-functional",
            "presentation", "negotiation", "decision making",
        ],
    },
    "Cover Letter": {
        "Structure": [
            "dear", "hiring manager", "position", "role", "apply",
            "interest", "company", "team", "opportunity",
        ],
        "Content Quality": [
            "experience", "skills", "project", "achievement",
            "contribution", "value", "passion", "motivated",
            "problem solving", "leadership",
        ],
        "Closing": [
            "thank you", "look forward", "interview", "discuss",
            "contact", "sincerely", "regards", "available",
        ],
    },
    "General / Other": {
        "Resume Basics": [
            "experience", "education", "skills", "objective", "summary",
            "contact", "email", "phone", "linkedin", "github",
        ],
        "Impact Words": [
            "achieved", "developed", "implemented", "managed", "led",
            "improved", "reduced", "increased", "designed", "built",
        ],
        "Credentials": [
            "bachelor", "master", "degree", "university", "certification",
            "gpa", "honors", "award",
        ],
    },
}

# Known companies to detect in resume content
KNOWN_COMPANIES = [
    "Google", "Amazon", "Meta", "Facebook", "Apple", "Microsoft", "Netflix",
    "TikTok", "ByteDance", "Stripe", "Shopify", "Uber", "Airbnb", "Twitter",
    "LinkedIn", "Salesforce", "Adobe", "Oracle", "IBM", "Intel", "Nvidia",
    "Tesla", "SpaceX", "Palantir", "Snap", "Pinterest", "DoorDash", "Lyft",
    "Robinhood", "Coinbase", "Square", "Block", "Twitch", "Reddit", "Discord",
    "Slack", "Zoom", "Atlassian", "GitHub", "GitLab", "MongoDB", "Snowflake",
    "Databricks", "Figma", "Notion", "Canva", "Vercel",
    # Canadian / local
    "Ubisoft", "UbiLab", "VistaVu", "YMCA", "Shaw", "Telus", "Rogers",
    "TD Bank", "RBC", "BMO", "Scotiabank", "CIBC", "Deloitte", "KPMG",
    "EY", "PwC", "Accenture", "CGI", "SAP", "BlackBerry",
    "University of Calgary", "University of Waterloo",
    # From filenames we saw
    "Stripe", "DoorDash", "Heli", "BWC",
]


# ═══════════════════════════════════════════════════════════════════════════════
# TEXT EXTRACTION
# ═══════════════════════════════════════════════════════════════════════════════

def extract_text_pdf(filepath):
    if not HAS_PYPDF2:
        return ""
    try:
        reader = PdfReader(str(filepath))
        text = ""
        for page in reader.pages[:10]:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text
    except Exception:
        return ""


def extract_text_docx(filepath):
    if not HAS_DOCX:
        return ""
    try:
        doc = DocxDocument(str(filepath))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        return ""


def extract_text_txt(filepath):
    try:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read(100_000)
    except Exception:
        return ""


def extract_text(filepath):
    ext = Path(filepath).suffix.lower()
    if ext == ".pdf":
        return extract_text_pdf(filepath)
    elif ext == ".docx":
        return extract_text_docx(filepath)
    elif ext in (".txt", ".rtf"):
        return extract_text_txt(filepath)
    return ""


# ═══════════════════════════════════════════════════════════════════════════════
# ANALYSIS ENGINE
# ═══════════════════════════════════════════════════════════════════════════════

def detect_role(filename, text):
    """Detect the role this resume/cover letter targets."""
    combined = (filename + " " + text).lower()

    # Check for cover letter first
    if "cover letter" in filename.lower() or "cover_letter" in filename.lower():
        return "Cover Letter"

    # Check for specific roles in filename or content
    role_signals = {
        "SDET": ["sdet", "software development engineer in test", "qa engineer", "test engineer", "quality assurance engineer"],
        "Data Scientist": ["data scientist", "data science", "machine learning engineer", "ml engineer", "data analyst"],
        "Web Developer": ["web developer", "frontend developer", "front-end developer", "web dev", "webdev"],
        "Product Manager": ["product manager", "program manager", "technical program manager"],
        "Software Engineer": ["software engineer", "sde", "swe", "backend engineer", "full stack engineer", "software developer"],
    }

    for role, signals in role_signals.items():
        for signal in signals:
            if signal in combined:
                return role

    # Fallback: check content keyword density
    if text:
        text_lower = text.lower()
        role_scores = {}
        for role, categories in ROLE_KEYWORDS.items():
            if role in ("General / Other", "Cover Letter", "Software Developer"):
                continue
            hits = 0
            total = 0
            for cat_keywords in categories.values():
                for kw in cat_keywords:
                    total += 1
                    if kw in text_lower:
                        hits += 1
            if total > 0:
                role_scores[role] = hits / total

        if role_scores:
            best_role = max(role_scores, key=role_scores.get)
            if role_scores[best_role] > 0.1:
                return best_role

    return "General / Other"


def detect_companies(text, filename):
    """Find company names mentioned in the resume."""
    combined = text + " " + filename
    found = []
    for company in KNOWN_COMPANIES:
        # Word-boundary search to avoid false positives
        pattern = re.compile(r'\b' + re.escape(company) + r'\b', re.IGNORECASE)
        if pattern.search(combined):
            found.append(company)

    # Also check for company names in the filename (e.g., "Rifat-Resume-Stripe-Internship")
    # Already covered by the above

    return list(set(found)) if found else ["None detected"]


def score_resume(text, role, filename):
    """
    Score a resume 0-100 as an expert hiring manager would.

    Scoring rubric:
      - Keyword coverage for role (0-30 pts)
      - Quantified achievements / impact metrics (0-20 pts)
      - Structure & formatting signals (0-15 pts)
      - Education & credentials (0-15 pts)
      - Experience depth (0-10 pts)
      - Professional presentation (0-10 pts)
    """
    if not text or len(text.strip()) < 50:
        return 10, "Unable to extract content for scoring"

    text_lower = text.lower()
    score = 0
    notes = []

    # --- 1. Keyword coverage for the role (0-30) ---
    role_kws = ROLE_KEYWORDS.get(role, ROLE_KEYWORDS["General / Other"])
    all_kws = []
    for cat_keywords in role_kws.values():
        all_kws.extend(cat_keywords)
    hits = sum(1 for kw in all_kws if kw in text_lower)
    kw_ratio = hits / max(len(all_kws), 1)
    kw_score = min(int(kw_ratio * 60), 30)  # up to 30
    score += kw_score
    if kw_ratio > 0.3:
        notes.append(f"Strong keyword coverage ({hits}/{len(all_kws)})")
    elif kw_ratio > 0.15:
        notes.append(f"Moderate keyword coverage ({hits}/{len(all_kws)})")
    else:
        notes.append(f"Weak keyword coverage ({hits}/{len(all_kws)})")

    # --- 2. Quantified achievements (0-20) ---
    # Look for numbers with context (%, $, x, etc.)
    metrics_patterns = [
        r'\d+%', r'\$[\d,]+', r'\d+x\b', r'\d+\+?\s*(users|customers|clients)',
        r'reduced\s+\w+\s+by', r'increased\s+\w+\s+by', r'improved\s+\w+\s+by',
        r'saved\s+\w+', r'generated\s+\w+', r'managed\s+\w+\s+team',
        r'\d+\s*projects?', r'led\s+\w+\s+of\s+\d+',
    ]
    metric_hits = sum(1 for p in metrics_patterns if re.search(p, text_lower))
    metrics_score = min(metric_hits * 4, 20)
    score += metrics_score
    if metrics_score >= 12:
        notes.append("Excellent use of quantified metrics")
    elif metrics_score >= 6:
        notes.append("Some quantified achievements")
    else:
        notes.append("Needs more quantified impact metrics")

    # --- 3. Structure & formatting signals (0-15) ---
    structure_signals = [
        "experience", "education", "skills", "projects", "summary",
        "objective", "work history", "technical skills", "achievements",
        "certifications", "awards", "publications", "volunteer",
    ]
    struct_hits = sum(1 for s in structure_signals if s in text_lower)
    struct_score = min(struct_hits * 3, 15)
    score += struct_score
    if struct_hits >= 5:
        notes.append("Well-structured sections")
    else:
        notes.append("Could improve section structure")

    # --- 4. Education & credentials (0-15) ---
    edu_signals = [
        "bachelor", "master", "mba", "phd", "b.sc", "m.sc", "b.eng", "m.eng",
        "university", "college", "degree", "gpa", "dean's list", "honors",
        "certification", "certified", "aws certified", "pmp",
    ]
    edu_hits = sum(1 for e in edu_signals if e in text_lower)
    edu_score = min(edu_hits * 3, 15)
    score += edu_score

    # --- 5. Experience depth (0-10) ---
    # Look for date ranges, job titles
    date_ranges = len(re.findall(r'20\d{2}\s*[-–]\s*(20\d{2}|present|current)', text_lower))
    exp_score = min(date_ranges * 3, 10)
    score += exp_score
    if date_ranges >= 3:
        notes.append(f"Strong experience ({date_ranges} roles)")

    # --- 6. Professional presentation (0-10) ---
    pro_signals = ["linkedin", "github", "portfolio", "email", "phone"]
    pro_hits = sum(1 for p in pro_signals if p in text_lower)
    pro_score = min(pro_hits * 2, 10)
    score += pro_score

    # Cap at 100
    score = min(score, 100)

    # Penalize very short resumes
    word_count = len(text.split())
    if word_count < 100:
        score = min(score, 35)
        notes.append(f"Very short ({word_count} words)")
    elif word_count < 200:
        score = max(score - 10, 0)
        notes.append(f"Somewhat short ({word_count} words)")

    return score, " | ".join(notes)


def analyze_keywords(text, role):
    """Return a dict of {category: [(keyword, found_bool), ...]} for the role."""
    role_kws = ROLE_KEYWORDS.get(role, ROLE_KEYWORDS["General / Other"])
    text_lower = text.lower() if text else ""

    results = {}
    for category, keywords in role_kws.items():
        results[category] = [(kw, kw in text_lower) for kw in keywords]
    return results


def is_actual_resume(row, text):
    """Filter out files that aren't actually resumes (e.g., game logs, license files)."""
    filename = row["filename"].lower()
    path = row["path"].lower()

    # Exclude obvious non-resumes
    skip_patterns = [
        "thirdpartylicense", "license", "log.txt", "build_log",
        "farming", "minecraft", "curseforge", "jre-legacy", "jre-x64",
        "case study", "essay", "module 3", "module 4",
        "power electronics", "converters",
        "change_of_enrolment", "medical_leave", "guarantor",
        "sublet", "volunteer agreement", "wochsublet",
        "abst participants", "study material",
        "individual licence",
        "visa statement",
        "employee contact",
        "1251_asg2", "asg3", "w25-657a",
        "ensf544 project", "predicting_price",
        "question_1_lean", "question_2_employee",
        "be 603", "be603",
        "a4.docx",
        "canadian canoe",
        "digital-transformation",
        "human resources policy",
        "student-programs-sde-transcript",
    ]
    for pattern in skip_patterns:
        if pattern in filename or pattern in path:
            return False

    # Must have some resume/cv signal in filename or decent confidence
    resume_name_signals = ["resume", "cv", "cover letter", "cover_letter"]
    has_name_signal = any(s in filename for s in resume_name_signals)

    if has_name_signal:
        return True

    # If confidence >= 60 and not filtered out, include
    try:
        if int(row.get("confidence", 0)) >= 60:
            return True
    except (ValueError, TypeError):
        pass

    return False


# ═══════════════════════════════════════════════════════════════════════════════
# EXCEL GENERATION
# ═══════════════════════════════════════════════════════════════════════════════

# Color palette
HEADER_FILL = PatternFill(start_color="1B1F3B", end_color="1B1F3B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(start_color="2D3250", end_color="2D3250", fill_type="solid")
SUBHEADER_FONT = Font(name="Calibri", bold=True, color="E8D5B7", size=10)

FOUND_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Green
FOUND_FONT = Font(name="Calibri", color="006100", size=10)
MISSING_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Red
MISSING_FONT = Font(name="Calibri", color="9C0006", size=10)

SCORE_HIGH_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
SCORE_MED_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
SCORE_LOW_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

ROW_ALT_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def create_excel(analyzed_resumes, output_path):
    """Create a beautifully formatted Excel spreadsheet."""
    wb = Workbook()

    # ── Sheet 1: Resume Overview ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resume Overview"

    headers = [
        "Resume Name", "Location (Folder)", "File Type", "Size (KB)",
        "Last Modified", "Tagged Role", "Company Mentions",
        "HM Score (0-100)", "Score Notes",
        "Keywords Found", "Keywords Missing", "Keyword Coverage %",
    ]

    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    # Write data rows
    for row_idx, resume in enumerate(analyzed_resumes, 2):
        row_data = [
            resume["filename"],
            resume["folder"],
            resume["extension"],
            resume["size_kb"],
            resume["modified"],
            resume["role"],
            ", ".join(resume["companies"]),
            resume["score"],
            resume["score_notes"],
            ", ".join(resume["found_keywords"]),
            ", ".join(resume["missing_keywords"]),
            resume["keyword_coverage_pct"],
        ]

        fill = ROW_ALT_FILL if row_idx % 2 == 0 else None

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill

        # Color-code the score cell
        score_cell = ws.cell(row=row_idx, column=8)
        if resume["score"] >= 70:
            score_cell.fill = SCORE_HIGH_FILL
            score_cell.font = Font(name="Calibri", bold=True, color="006100", size=11)
        elif resume["score"] >= 45:
            score_cell.fill = SCORE_MED_FILL
            score_cell.font = Font(name="Calibri", bold=True, color="9C5700", size=11)
        else:
            score_cell.fill = SCORE_LOW_FILL
            score_cell.font = Font(name="Calibri", bold=True, color="9C0006", size=11)

        # Color-code coverage %
        cov_cell = ws.cell(row=row_idx, column=12)
        pct = resume["keyword_coverage_pct"]
        if pct >= 35:
            cov_cell.fill = SCORE_HIGH_FILL
        elif pct >= 20:
            cov_cell.fill = SCORE_MED_FILL
        else:
            cov_cell.fill = SCORE_LOW_FILL

    # Column widths
    col_widths = [35, 55, 8, 10, 16, 22, 30, 14, 45, 55, 55, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Freeze the header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(analyzed_resumes) + 1}"

    # ── Sheet 2: Keyword Analysis (detailed per resume) ──────────────────────
    ws2 = wb.create_sheet("Keyword Analysis")

    kw_headers = ["Resume Name", "Role", "Keyword Category", "Keyword", "Status"]
    for col_idx, header in enumerate(kw_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    kw_row = 2
    for resume in analyzed_resumes:
        for category, kw_results in resume["keyword_analysis"].items():
            for keyword, found in kw_results:
                ws2.cell(row=kw_row, column=1, value=resume["filename"]).font = Font(name="Calibri", size=10)
                ws2.cell(row=kw_row, column=2, value=resume["role"]).font = Font(name="Calibri", size=10)
                ws2.cell(row=kw_row, column=3, value=category).font = Font(name="Calibri", size=10, bold=True)

                kw_cell = ws2.cell(row=kw_row, column=4, value=keyword)
                status_cell = ws2.cell(row=kw_row, column=5)

                if found:
                    kw_cell.font = FOUND_FONT
                    kw_cell.fill = FOUND_FILL
                    status_cell.value = "FOUND"
                    status_cell.font = FOUND_FONT
                    status_cell.fill = FOUND_FILL
                else:
                    kw_cell.font = MISSING_FONT
                    kw_cell.fill = MISSING_FILL
                    status_cell.value = "MISSING"
                    status_cell.font = MISSING_FONT
                    status_cell.fill = MISSING_FILL

                for c in range(1, 6):
                    ws2.cell(row=kw_row, column=c).border = THIN_BORDER
                    ws2.cell(row=kw_row, column=c).alignment = Alignment(vertical="top")

                kw_row += 1

    # Column widths for sheet 2
    kw_col_widths = [35, 22, 22, 25, 12]
    for i, width in enumerate(kw_col_widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = width

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:E{kw_row - 1}"

    # ── Sheet 3: Summary Statistics ──────────────────────────────────────────
    ws3 = wb.create_sheet("Summary")

    # Title
    title_cell = ws3.cell(row=1, column=1, value="Resume Analysis Summary")
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="1B1F3B")
    ws3.merge_cells("A1:D1")

    # Stats
    stats = [
        ("Total Resumes Analyzed", len(analyzed_resumes)),
        ("Average HM Score", round(sum(r["score"] for r in analyzed_resumes) / max(len(analyzed_resumes), 1), 1)),
        ("Highest Score", max((r["score"] for r in analyzed_resumes), default=0)),
        ("Lowest Score", min((r["score"] for r in analyzed_resumes), default=0)),
        ("", ""),
        ("By Role:", ""),
    ]

    # Count by role
    role_counts = defaultdict(int)
    role_avg_scores = defaultdict(list)
    for r in analyzed_resumes:
        role_counts[r["role"]] += 1
        role_avg_scores[r["role"]].append(r["score"])

    for role, count in sorted(role_counts.items(), key=lambda x: -x[1]):
        avg = round(sum(role_avg_scores[role]) / max(count, 1), 1)
        stats.append((f"  {role}", f"{count} files (avg score: {avg})"))

    for i, (label, value) in enumerate(stats, 3):
        label_cell = ws3.cell(row=i, column=1, value=label)
        label_cell.font = Font(name="Calibri", bold=True, size=11)
        val_cell = ws3.cell(row=i, column=2, value=value)
        val_cell.font = Font(name="Calibri", size=11)

    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 35

    # ── Save ─────────────────────────────────────────────────────────────────
    wb.save(output_path)
    print(f"  Excel saved to: {output_path}")


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════

def main():
    csv_path = Path(__file__).parent / "resumes_found.csv"
    output_path = Path(__file__).parent / "Resume_Analysis.xlsx"

    if not csv_path.exists():
        print("ERROR: resumes_found.csv not found. Run resume_finder.py --deep first.")
        sys.exit(1)

    # Load CSV
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        all_rows = list(reader)

    print(f"  Loaded {len(all_rows)} entries from CSV")

    # Filter to actual resumes/cover letters
    analyzed = []
    for i, row in enumerate(all_rows):
        filepath = row["path"]
        if not os.path.exists(filepath):
            continue

        print(f"  [{i+1}/{len(all_rows)}] Analyzing: {row['filename']}...", end="\r")

        # Extract text
        text = extract_text(filepath)

        # Filter out non-resumes
        if not is_actual_resume(row, text):
            continue

        # Detect role
        role = detect_role(row["filename"], text)

        # Detect companies
        companies = detect_companies(text, row["filename"])

        # Score
        score, score_notes = score_resume(text, role, row["filename"])

        # Keyword analysis
        kw_analysis = analyze_keywords(text, role)

        # Flatten keywords for overview sheet
        found_kws = []
        missing_kws = []
        total_kws = 0
        for cat_results in kw_analysis.values():
            for kw, found in cat_results:
                total_kws += 1
                if found:
                    found_kws.append(kw)
                else:
                    missing_kws.append(kw)

        coverage_pct = round((len(found_kws) / max(total_kws, 1)) * 100, 1)

        analyzed.append({
            "filename": row["filename"],
            "folder": row["folder"],
            "extension": row["extension"],
            "size_kb": float(row["size_kb"]),
            "modified": row["modified"],
            "role": role,
            "companies": companies,
            "score": score,
            "score_notes": score_notes,
            "found_keywords": found_kws,
            "missing_keywords": missing_kws,
            "keyword_coverage_pct": coverage_pct,
            "keyword_analysis": kw_analysis,
        })

    print(f"\n  Filtered to {len(analyzed)} actual resumes/cover letters")

    # Sort by score descending
    analyzed.sort(key=lambda x: x["score"], reverse=True)

    # Generate Excel
    create_excel(analyzed, output_path)
    print(f"\n  Done! {len(analyzed)} resumes analyzed.")
    print(f"  Output: {output_path}")


if __name__ == "__main__":
    main()
